#!/usr/bin/env python3
"""生成したスライドのレイアウト検査。check.py（ファイル構造の検査）と併用する。

    python3 lint.py deck.pptx   （install.sh を使ったなら <skill>/.venv/bin/python）

本人の資料の実測値（本文の下端の中央値 6.8in、左右 0.54〜12.66in、
1枚あたり本文 250〜300 字＋図形 10〜15 個）から外れるものを警告する。
警告は「直すか、直さない理由を言えるか」で判断する。ERROR は必ず直す。
"""
import re
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

E = 914400
SW, SH = 13.333, 7.5
SIZES = {12, 14, 16, 18, 20, 22, 24, 28, 36}
LATEX_LEAK = re.compile(r"\$|\\[a-zA-Z]+|[A-Za-z]_\{|[A-Za-z]\^\{|\b(mu|sigma|Delta|delta)_[a-z0-9]")


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)


def bbox(sh):
    x, y, w, h = sh.left / E, sh.top / E, sh.width / E, sh.height / E
    if round(getattr(sh, "rotation", 0) or 0) % 180 == 90:
        # 90°/270° 回転した図形は、中心を保ったまま縦横が入れ替わって見える
        cx, cy = x + w / 2, y + h / 2
        w, h = h, w
        x, y = cx - w / 2, cy - h / 2
    return x, y, x + w, y + h


def lint(path):
    p = Presentation(path)
    n_err = n_warn = 0
    for i, s in enumerate(p.slides, 1):
        msgs = []
        xml = s._element.xml
        divider = "<p:bg>" in xml and "16223B" in xml
        shapes = [sh for sh in walk(s.shapes) if sh.left is not None and not sh.is_placeholder]
        texts = [sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title = next((sh.text_frame.text for sh in texts if sh.name.startswith("Title")), "")
        is_cover = any("発表資料" in sh.text_frame.text for sh in texts)
        is_toc = title.startswith("Table of Contents") or title.startswith("目次")
        # 1. キャンバス外
        for sh in shapes:
            x0, y0, x1, y1 = bbox(sh)
            if x0 < -0.01 or y0 < -0.01 or x1 > SW + 0.01 or y1 > SH + 0.01:
                msgs.append(("ERROR", "キャンバス外: %s (%.2f,%.2f)-(%.2f,%.2f)" % (sh.name, x0, y0, x1, y1)))
        # 2. 文字サイズと数式の書き漏れ
        for sh in texts:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.font.size and r.font.size.pt not in SIZES and r.text.strip():
                        msgs.append(("WARN", "規定外の文字サイズ %.1fpt: %r" % (r.font.size.pt, r.text[:30])))
                    if r.font.size and r.font.size.pt < 12 and r.text.strip():
                        msgs.append(("ERROR", "12pt 未満: %r" % r.text[:30]))
            t = sh.text_frame.text
            if LATEX_LEAK.search(t):
                msgs.append(("ERROR", "数式が文字のまま（数式モードにする）: %r" % t[:60]))
        if not (divider or is_cover or is_toc):
            body = [sh for sh in shapes if sh.top / E > 1.2 and sh.name != "Slide Number"]
            if body:
                bottom = max(bbox(sh)[3] for sh in body)
                right = max(bbox(sh)[2] for sh in body)
                left = min(bbox(sh)[0] for sh in body)
                if bottom < 5.6:
                    msgs.append(("WARN", "本文が %.2fin で終わる（本人の資料は 6.8in 前後まで使う）。下が空きすぎ" % bottom))
                if right < 11.3:
                    msgs.append(("WARN", "本文の右端が %.2fin（右が空いている。図を大きくするか幅を広げる）" % right))
                if left > 1.2:
                    msgs.append(("WARN", "本文の左端が %.2fin（左余白は 0.54in に揃える）" % left))
            chars = sum(len(sh.text_frame.text) for sh in texts if not sh.name.startswith("Title"))
            n_math = xml.count("<m:oMath")
            figs = [sh for sh in shapes if not sh.has_text_frame or not sh.text_frame.text.strip()]
            if chars < 60 and n_math == 0 and len(figs) < 3:
                msgs.append(("WARN", "情報が少ない（本文 %d 字、図形 %d 個）。他のページと統合を検討" % (chars, len(figs))))
            # 3. 文章だけの2カラム
            #    points()（色つき見出しの対比の 2 列、本人の修正版の型）は除く
            cols = [sh for sh in texts if sh.width / E < 6.3 and sh.top / E > 2.2 and len(sh.text_frame.text) > 60
                    and not sh.name.startswith("Points")]
            lefts = [sh for sh in cols if sh.left / E < 4]
            rights = [sh for sh in cols if sh.left / E > 6]
            if lefts and rights and len(figs) < 3:
                msgs.append(("WARN", "文章だけの2カラム。1カラムにするか、片側を図にする"))
            # 4. テキストボックス同士の重なり（見積もり高さで判定するので目視で確認）
            for a in range(len(texts)):
                for b in range(a + 1, len(texts)):
                    A, B = bbox(texts[a]), bbox(texts[b])
                    ox = min(A[2], B[2]) - max(A[0], B[0])
                    oy = min(A[3], B[3]) - max(A[1], B[1])
                    if texts[a].name.startswith("FieldsetLabel") or texts[b].name.startswith("FieldsetLabel"):
                        continue        # fieldset() の見出しは枠の上辺に重ねる作り
                    if ox > 0.3 and oy > 0.12:
                        msgs.append(("WARN", "文字の重なりの疑い: %r と %r" % (texts[a].text_frame.text[:20], texts[b].text_frame.text[:20])))
            # 6. 箱の使いすぎ（Claude の癖。本人の修正 2026-10-15：箇条書きは箱で囲まない）
            filled = []
            for sh in texts:
                if sh.name.startswith("Title") or abs(sh.top / E - 1.27) < 0.02:
                    continue            # タイトルとリード帯
                try:
                    if sh.fill.type != MSO_FILL.SOLID:
                        continue
                    rgb = str(sh.fill.fore_color.rgb).upper()
                except Exception:
                    continue
                if rgb in ("FFFFFF", "FFFEFF"):
                    continue            # 白地（擬似コードの箱・見出しの下地）
                filled.append(sh)
                paras = [pp for pp in sh.text_frame.paragraphs if "".join(r.text for r in pp.runs).strip()]
                bul = sum(1 for pp in paras if pp._p.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None)
                if bul >= 2 or len(paras) >= 3:
                    msgs.append(("WARN", "箇条書き・複数段落を塗りの箱に入れている: %r（箱で囲まず、色つき見出し＋"
                                         "points() にする）" % sh.text_frame.text[:30]))
            if len(filled) >= 3:
                msgs.append(("WARN", "文字入りの塗りの箱が %d 個（1 枚に 2 個まで。見出し＋箇条書きか図にする）" % len(filled)))
            # 5. ノート
            if not (s.has_notes_slide and s.notes_slide.notes_text_frame is not None
                    and s.notes_slide.notes_text_frame.text.strip()):
                msgs.append(("WARN", "ノートが空（論文ゼミはタイトルと本文の日本語訳を入れる）"))
        if msgs:
            print("--- slide %d  %s" % (i, title[:60]))
            for lv, m in msgs:
                print("  [%s] %s" % (lv, m))
                n_err += lv == "ERROR"
                n_warn += lv == "WARN"
    print("\n%d slides, %d ERROR, %d WARN" % (len(p.slides), n_err, n_warn))
    return n_err


if __name__ == "__main__":
    sys.exit(1 if lint(sys.argv[1]) else 0)
