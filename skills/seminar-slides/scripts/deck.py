#!/usr/bin/env python3
"""泉 知成のゼミ・発表スライドを、本人の過去資料と同じデザインで組むエンジン。

土台は assets/base.pptx（論文ゼミ_泉0421.pptx から全スライドを抜いたもの）。
スライドマスター・レイアウト（タイトル左の紺の縦線・ページ番号）・テーマは
そのまま継承される。ここでは本文の図形だけを、実測した寸法と色で置く。

    import sys; sys.path.insert(0, "<skill>/scripts")
    from deck import Deck, C
    d = Deck()
    d.cover(tag="論文ゼミ 発表資料", title="...", authors="A¹, B²",
            venue="arXiv (2021)", date="05.10.2026", grade="M2", name="泉 知成",
            affiliations=["¹ ...", "² ..."])
    s = d.slide("Setup: Regret", lead="**Regret** measures ...")
    y = s.text(C.L, s.top, C.W, "Let $\\mu_k$ be ...")
    y = s.eq(r"R(T) = ...", y + 0.1, number="(1)")
    s.notes("Setup: 後悔\\n\\n本文の日本語訳 ...")
    d.save("out.pptx")

python-pptx と lxml が要る（リポジトリの install.sh が .venv に入れる。スキルの .venv からも使える）。
"""
import math
import os
import re
import sys
import unicodedata

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import omml  # noqa: E402

ASSETS = os.path.join(HERE, "..", "assets")
BASE = os.path.join(ASSETS, "base.pptx")
ICON = {k: os.path.join(ASSETS, "icons", k + ".png") for k in ("human", "slot", "lock")}
# 本人の過去資料（paste() の複製元）。短い名前でも、ファイル名だけでも指定できる
REF_DIR = os.path.join(ASSETS, "reference")
REF = {"paper0421": os.path.join(REF_DIR, "論文ゼミ_泉0421.pptx"),
       "thesis": os.path.join(REF_DIR, "泉知成_修論中間発表スライド_完成.pptx"),
       "fixes": os.path.join(REF_DIR, "スライドデザイン変更点まとめ.pptx")}
EMU = 914400
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


# --- 色（本人の資料からの実測値。勝手に濃くしたり変えたりしない） ---------------------
PAL = {
    "ink": "0D114A",        # タイトル・紺の図形・矢印・強調語
    "body": "1C1F28",       # 本文
    "navy-bg": "16223B",    # セクション扉の背景・目次の文字
    "warm": "32302A",       # 表紙の日付・縦線・脚注・図中ラベル
    "red": "D14B46",        # 衝突・注意・課題・表紙タグ
    "green": "196723",      # 段階1（初期化など）・肯定
    "blue": "0070C0",       # 段階2・引用 [Author+, Year]
    "gray": "A6A6A6",       # 現在地以外のグレーアウト
    "line-gray": "BFBFBF",  # 擬似コード枠・補助線
    "lead": "D9DBE8",       # リード文の帯・角丸の箱（2026-10-14 本人の指定。旧 EEF1F9）
    "pale-blue": "D1E1EA",  # 擬似コードの現在行・2つ目の面
    "yellow": "F6EAC3",     # 数式・表のハイライト（2026-10-14 本人の指定。旧 FFF1C8）
    "amber": "FFC000",      # ハイライト枠
    "pale-green": "E4F6DD", # 緑系の面
    "pale-red": "FFCEC6",   # 赤系の面（課題）（2026-10-14 本人の指定。旧 FDECEC）
    "salmon": "FFC6BE",     # 衝突の爆発図形
    "white": "FFFFFF",
    "bg": "FFFEFF",
}
FONT_LATIN = "Segoe UI"
FONT_EA = "Hiragino Sans W4"      # ヒラギノ角ゴシック W4


def col(c):
    return PAL.get(c, c)


# --- グリッド（インチ。実測値） ------------------------------------------------------
class C:
    SW, SH = 13.333, 7.5
    L = 0.54            # 本文の左端（リード帯の左端）
    R = 12.66           # 本文の右端
    W = R - L           # 12.12
    TITLE_X, TITLE_Y, TITLE_H = 0.67, 0.24, 1.03
    LEAD_Y = 1.27
    TOP_NOLEAD = 1.40   # リード帯がない時の本文の開始
    FLOOR = 6.85        # 本文の下限（ページ番号 6.95 に触れない）
    LINE20 = 0.36       # 20pt 1行の送り（PowerPoint の Segoe UI の実測。0.333 では下に積むほどずれる）
    PAD = 0.10          # テキストボックス上下の内側余白の合計


# --- 文字幅の見積もり（積み上げ用） ----------------------------------------------------
def _cells(s):
    n = 0.0
    for ch in s:
        if unicodedata.east_asian_width(ch) in ("W", "F"):
            n += 1.9
        elif ch in "il.,;:'|!()[] ":
            n += 0.55
        elif ch.isupper() or ch in "mwMW":
            n += 1.25
        else:
            n += 1.0
    return n


def _plain(markup):
    s = re.sub(r"\$(.+?)\$", lambda m: "x" * max(1, int(len(m.group(1)) * 0.45)), str(markup))
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"\{[\w-]+:(.+?)\}", r"\1", s)
    return s


def lines_of(markup, w, size=20, indent=0.0):
    """幅 w インチに size pt で組んだときの行数の見積もり。"""
    per = (w - 0.2 - indent) / (0.140 * size / 20.0)
    n = 0
    for p in str(markup).split("\n"):
        n += max(1, math.ceil(_cells(_plain(p)) / max(per, 4)))
    return n


def text_h(markup, w, size=20, spacing=1.0, indent=0.0):
    return lines_of(markup, w, size, indent) * C.LINE20 * size / 20.0 * spacing + C.PAD


# --- 文字列 → DrawingML ------------------------------------------------------------
RICH = re.compile(r"\*\*(.+?)\*\*|\{([\w-]+):(.+?)\}|\$(.+?)\$", re.S)


def esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def rpr(size=20, color="body", bold=False, italic=False, underline=False):
    return ('<a:rPr lang="en-US" altLang="ja-JP" sz="%d" b="%d" i="%d"%s dirty="0">'
            '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
            '<a:latin typeface="%s"/><a:ea typeface="%s"/><a:cs typeface="%s"/></a:rPr>'
            % (round(size * 100), 1 if bold else 0, 1 if italic else 0,
               ' u="sng"' if underline else "", col(color), FONT_LATIN, FONT_EA, FONT_LATIN))


def _r(text, size, color, bold=False, italic=False, underline=False):
    return "<a:r>%s<a:t>%s</a:t></a:r>" % (rpr(size, color, bold, italic, underline), esc(text))


def runs(markup, size=20, color="body", bold=False, italic=False):
    """記法: **太字** / {red:色付き太字} / {red-n:色付き細字} / {i:斜体} / {u:下線} / $数式$"""
    out, pos, s = [], 0, str(markup)
    for m in RICH.finditer(s):
        if m.start() > pos:
            out.append(_r(s[pos:m.start()], size, color, bold, italic))
        # 太字・色の中に $数式$ があってもよいように、中身をもう一度 runs() に通す
        if m.group(1) is not None:
            out.append(runs(m.group(1), size, color, True, italic))
        elif m.group(2) is not None:
            tok, txt = m.group(2), m.group(3)
            if tok == "u":
                out.append(_r(txt, size, color, bold, italic, True))
            elif tok == "i":
                out.append(_r(txt, size, color, bold, True))
            elif tok.endswith("-n"):
                out.append(runs(txt, size, tok[:-2], bold, italic))
            else:
                out.append(runs(txt, size, tok, True, italic))
        else:
            out.append(omml.inline_math(m.group(4), size, col(color)))
        pos = m.end()
    if pos < len(s):
        out.append(_r(s[pos:], size, color, bold, italic))
    if not out:
        out.append('<a:endParaRPr lang="en-US" sz="%d"/>' % round(size * 100))
    return "".join(out)


def para(markup, size=20, color="body", bold=False, align="l", bullet=None,
         spacing=None, before=0, indent=0.0, italic=False, after=0):
    """1 段落。bullet="•" / "1." / "-"。折り返した行は記号の右に揃う。
    行頭の記号の色は、段落の最初の文字の色になる（PowerPoint の既定）。"""
    sp = ""
    if spacing:
        sp += '<a:lnSpc><a:spcPct val="%d"/></a:lnSpc>' % int(spacing * 100000)
    if before:
        sp += '<a:spcBef><a:spcPts val="%d"/></a:spcBef>' % int(before * 100)
    if after:
        sp += '<a:spcAft><a:spcPts val="%d"/></a:spcAft>' % int(after * 100)
    m = int(indent * EMU)
    if bullet == "•":
        ppr = ('<a:pPr marL="%d" indent="-285750" algn="%s">%s<a:buFont typeface="Arial"/>'
               '<a:buChar char="&#8226;"/></a:pPr>' % (285750 + m, align, sp))
    elif bullet == "1.":
        ppr = ('<a:pPr marL="%d" indent="-342900" algn="%s">%s<a:buFont typeface="+mj-lt"/>'
               '<a:buAutoNum type="arabicPeriod"/></a:pPr>' % (342900 + m, align, sp))
    elif bullet == "-":
        ppr = ('<a:pPr marL="%d" indent="-228600" algn="%s">%s<a:buFont typeface="Arial"/>'
               '<a:buChar char="&#8211;"/></a:pPr>' % (571500 + m, align, sp))
    else:
        ppr = '<a:pPr marL="%d" indent="0" algn="%s">%s<a:buNone/></a:pPr>' % (m, align, sp)
    return "<a:p>%s%s</a:p>" % (ppr, runs(markup, size, color, bold, italic))


NOTE_COLOR = "222224"


def _strip_transitions(xml):
    """<p:transition>（変形 morph を包む mc:AlternateContent を含む）をすべて外す。"""
    xml = re.sub(r"<mc:AlternateContent[^>]*>\s*<mc:Choice[^>]*>\s*<p:transition.*?</mc:AlternateContent>", "", xml, flags=re.S)
    xml = re.sub(r"<p:transition\b[^>]*/>", "", xml)
    return re.sub(r"<p:transition\b.*?</p:transition>", "", xml, flags=re.S)


def note_para(line, first=False, size=11):
    """ノートの 1 段落（本人のノートの書式：11pt、色 222224、箇条書きは • で marL 171450）。"""
    lvl = 0
    s = line
    if s.startswith("  - "):
        lvl, s = 2, s[4:]
    elif s.startswith("・") or s.startswith("- "):
        lvl, s = 1, s[1:] if s.startswith("・") else s[2:]
    bold_all = first          # タイトル行は全体を太字（数式を含んでもよいように記法では囲まない）
    if not first and s.startswith("【") and "】" in s[:12]:
        s = "**%s**%s" % (s[:s.index("】") + 1], s[s.index("】") + 1:])
    if lvl:
        mar = 171450 if lvl == 1 else 400050
        ppr = ('<a:pPr marL="%d" indent="-171450"><a:buFont typeface="Arial"/><a:buChar char="%s"/></a:pPr>'
               % (mar, "&#8226;" if lvl == 1 else "&#8211;"))
    else:
        ppr = '<a:pPr marL="0" indent="0"><a:buNone/></a:pPr>'
    body = runs(s.replace("**", "") if bold_all else s, size, NOTE_COLOR, bold=bold_all) if s.strip() \
        else '<a:endParaRPr lang="ja-JP" sz="%d"/>' % (size * 100)
    return "<a:p>%s%s</a:p>" % (ppr, body)


def _math_as_text(m):
    """数式ブロックを、同じ大きさ・色の普通の文字に置き換える（表の代替表示用）。"""
    blk = m.group(0)
    txt = "".join(re.findall(r"<m:t[^>]*>(.*?)</m:t>", blk, flags=re.S))
    sz = re.search(r'sz="(\d+)"', blk)
    c = re.search(r'srgbClr val="([0-9A-Fa-f]{6})"', blk)
    return ('<a:r><a:rPr lang="en-US" sz="%s" i="1" dirty="0"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
            '<a:latin typeface="Cambria Math"/></a:rPr><a:t>%s</a:t></a:r>'
            % (sz.group(1) if sz else "1600", c.group(1) if c else PAL["body"], txt))


def has_math(xml):
    return "<a14:m" in xml


def _frag(xml):
    return etree.fromstring('<root xmlns:a="%s" xmlns:p="%s" xmlns:r="%s" xmlns:a14="%s" xmlns:mc="%s">%s</root>'
                            % (NS["a"], NS["p"], NS["r"], A14, MC, xml))[0]


# --- スライド --------------------------------------------------------------------
class Slide:
    def __init__(self, deck, slide, top, kind="body"):
        self.deck, self.s, self.top, self.kind = deck, slide, top, kind
        self.tree = slide.shapes._spTree
        self.log = []   # (kind, x, y, w, h)  検査用

    # ---- 低レベル ----
    def _id(self):
        ids = [int(e.get("id")) for e in self.tree.iter("{%s}cNvPr" % NS["p"]) if e.get("id", "").isdigit()]
        return max(ids + [100]) + 1

    def _append(self, xml, back=False):
        el = _frag(xml)
        if back:
            self.tree.insert(2, el)
        else:
            self.tree.append(el)
        return el

    @staticmethod
    def _wrap_math(sp_xml):
        """数式（a14:m）は包まずにそのまま置く。PowerPoint は読み込んで数式として扱い、
        保存時に自分で AlternateContent を付ける。包むと Keynote が代替側（数式なし）を描く。"""
        return sp_xml

    def _sp(self, x, y, w, h, prst="rect", fill=None, line=None, lw=1.0, dash=None,
            body="", anchor="t", ins=(0.1, 0.05, 0.1, 0.05), txbox=False, autofit=False,
            adj=None, wrap=True, back=False, rot=0, flip="", tag=None):
        sid = self._id()
        av = "".join('<a:gd name="%s" fmla="val %d"/>' % (k, v) for k, v in (adj or {}).items())
        fill_xml = '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % col(fill) if fill else "<a:noFill/>"
        if line:
            ln = '<a:ln w="%d"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>%s</a:ln>' % (
                int(lw * 12700), col(line), '<a:prstDash val="%s"/>' % dash if dash else "")
        else:
            ln = "<a:ln><a:noFill/></a:ln>"
        bp = ('<a:bodyPr wrap="%s" lIns="%d" tIns="%d" rIns="%d" bIns="%d" rtlCol="0" anchor="%s">%s</a:bodyPr>'
              % ("square" if wrap else "none", *(int(v * EMU) for v in ins), anchor,
                 "<a:spAutoFit/>" if autofit else "<a:noAutofit/>"))
        if not body:
            body = '<a:p><a:endParaRPr lang="ja-JP"/></a:p>'
        xml = ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s %d"/><p:cNvSpPr%s/><p:nvPr/></p:nvSpPr>'
               '<p:spPr><a:xfrm%s%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
               '<a:prstGeom prst="%s"><a:avLst>%s</a:avLst></a:prstGeom>%s%s</p:spPr>'
               '<p:txBody>%s<a:lstStyle/>%s</p:txBody></p:sp>'
               % (sid, tag or ("TextBox" if txbox else prst), sid, ' txBox="1"' if txbox else "",
                  ' rot="%d"' % int(rot * 60000) if rot else "", flip,
                  int(x * EMU), int(y * EMU), int(w * EMU), int(h * EMU), prst, av, fill_xml, ln, bp, body))
        self._append(self._wrap_math(xml), back=back)
        self.log.append((tag or ("txt" if txbox else prst), x, y, w, h))

    # ---- 文字 ----
    def text(self, x, y, w, markup, size=20, color="body", bold=False, align="l",
             h=None, spacing=None, bullet=None, gap=0, anchor="t", italic=False, indent=0.0, tag=None):
        """テキストボックス。改行で段落を分ける。戻り値は下端の y（見積もり）。
        bullet="•" のとき各段落が箇条書き。gap は段落間の pt。"""
        paras = str(markup).split("\n")
        body = "".join(para(p, size, color, bold, align, bullet, spacing,
                            before=(gap if i else 0), indent=indent, italic=italic)
                       for i, p in enumerate(paras))
        ind = (0.3125 if bullet == "•" else 0.375 if bullet == "1." else 0.625 if bullet == "-" else 0.0) + indent
        est = text_h(markup, w, size, spacing or 1.0, ind) + gap / 72.0 * (len(paras) - 1)
        self._sp(x, y, w, h or est, body=body, txbox=True, autofit=h is None, anchor=anchor, tag=tag)
        return y + (h or est)

    def bullets(self, x, y, w, items, size=20, color="body", gap=6, bullet="•", spacing=None):
        return self.text(x, y, w, "\n".join(items), size, color, bullet=bullet, gap=gap, spacing=spacing)

    def label(self, x, y, markup, size=20, color="body", bold=False):
        """短いラベル（折り返さない）。左上指定。"""
        w = _cells(_plain(markup)) * 0.125 * size / 20 + 0.2
        h = C.LINE20 * size / 20 + 0.08
        self._sp(x, y, w, h, body=para(markup, size, color, bold), txbox=True, wrap=False,
                 autofit=True, ins=(0.05, 0.04, 0.05, 0.04))
        return x + w

    def clabel(self, cx, y, markup, size=20, color="body", bold=False, w=2.0):
        """中心 x を指定するラベル。"""
        return self.text(cx - w / 2, y, w, markup, size, color, bold, align="ctr")

    def heading(self, x, y, markup, color="body", size=20, w=None):
        """本文中の小見出し（"Example:" など。太字 20pt）。"""
        return self.text(x, y, w or (C.R - x), markup, size, color, bold=True)

    def num(self, x, y, n, color="ink", size=28):
        """❶❷❸ の丸数字。左上指定、幅 0.62。"""
        ch = "❶❷❸❹❺❻❼❽❾❿"[n - 1] if isinstance(n, int) else n
        self._sp(x, y, 0.62, 0.57, body=para(ch, size, color, align="ctr"), txbox=True)
        return y + 0.57

    def steps(self, x, y, w, items, color="ink", gap=0.62, size=20, start=1):
        """❶ 見出し: 説明 を縦に並べ、間に小さな三角を置く（本人の Overview スライドの型）。
        items は (見出し, 説明) のリスト。"""
        for i, (head, desc) in enumerate(items):
            yy = y + i * gap
            self.num(x, yy - 0.06, start + i, color)
            self.text(x + 0.62, yy, w - 0.62, "{%s:%s}  %s" % (color, head, desc) if head else desc, size)
            if i < len(items) - 1:
                self.shape("triangle", x + 0.24, yy + 0.53, 0.14, 0.12, fill=color, line=color, flipV=True)
        return y + len(items) * gap

    # ---- 数式 ----
    def eq(self, latex, y, x=None, w=None, size=22, number=None, color="body", align="ctr", h=None):
        """PowerPoint の数式オブジェクト（画像にしない）。number は式番号を同じ行の右端に。"""
        x = C.L if x is None else x
        w = (C.R - x - (1.1 if number else 0)) if w is None else w
        tall = any(k in latex for k in ("\\frac", "\\sum", "\\int", "\\prod", "\\begin", "_{\\substack",
                                        "\\underbrace", "\\overbrace", "\\max_", "\\min_", "\\arg"))
        h = h or (0.95 if tall else 0.55) * size / 22
        jc = {"ctr": "centerGroup", "l": "left", "r": "right"}[align]
        self._sp(x, y, w, h, body='<a:p><a:pPr algn="%s"/>%s</a:p>' % (align, omml.block_math(latex, size, col(color), jc)),
                 txbox=True, anchor="ctr", tag="Equation")
        if number:
            self._sp(C.R - 1.05, y, 1.05, h, body=para(number, 20, color, align="r"), txbox=True, anchor="ctr")
        return y + h

    # ---- 図形 ----
    def box(self, x, y, w, h, markup="", fill="lead", line=None, lw=1.0, prst="roundRect",
            size=20, color="body", bold=False, align="ctr", anchor="ctr", radius=0.10,
            dash=None, ins=(0.1, 0.05, 0.1, 0.05), spacing=None, bullet=None, gap=0, back=False):
        """角丸の箱など。radius は角の半径（インチ）。"""
        if h is None:   # 文字の高さに合わせる（上下に余分な空きを作らない）
            h = sum(text_h(p_, w - ins[0] - ins[2], size) - C.PAD for p_ in str(markup).split("\n")) + 0.18 + ins[1] + ins[3]
        adj = {"adj": int(min(0.5, radius / max(min(w, h), 0.05)) * 100000)} if prst == "roundRect" else None
        body = "".join(para(p, size, color, bold, align, bullet, spacing, before=(gap if i else 0))
                       for i, p in enumerate(str(markup).split("\n"))) if markup else ""
        self._sp(x, y, w, h, prst=prst, fill=fill, line=line, lw=lw, dash=dash, body=body,
                 anchor=anchor, adj=adj, ins=ins, back=back)
        return y + h

    def shape(self, prst, x, y, w, h, fill=None, line="ink", lw=1.0, adj=None, markup="",
              size=16, color="body", bold=False, rot=0, flipH=False, flipV=False, dash=None):
        """任意のプリセット図形（chevron, homePlate, rightArrow, downArrow, triangle,
        leftBrace, rightBrace, irregularSeal2, ellipse, ...）。"""
        body = para(markup, size, color, bold, "ctr") if markup else ""
        flip = (' flipH="1"' if flipH else "") + (' flipV="1"' if flipV else "")
        self._sp(x, y, w, h, prst=prst, fill=fill, line=line, lw=lw, adj=adj, body=body,
                 anchor="ctr", rot=rot, flip=flip, dash=dash, ins=(0.02, 0.02, 0.02, 0.02))
        return y + h

    def arrow(self, x1, y1, x2, y2, color="ink", lw=3.5, head="end", dash=None):
        """矢印。プレイヤー→腕の「引く」矢印は 3.5pt、図の流れは 1.5pt。"""
        sid = self._id()
        x, y, w, h = min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1)
        flip = (' flipH="1"' if x2 < x1 else "") + (' flipV="1"' if y2 < y1 else "")
        he = '<a:headEnd type="triangle"/>' if head in ("start", "both") else ""
        te = '<a:tailEnd type="triangle"/>' if head in ("end", "both") else ""
        xml = ('<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="%d" name="Arrow %d"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
               '<p:spPr><a:xfrm%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
               '<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>'
               '<a:ln w="%d"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>%s%s%s</a:ln></p:spPr></p:cxnSp>'
               % (sid, sid, flip, int(x * EMU), int(y * EMU), int(w * EMU), int(h * EMU),
                  int(lw * 12700), col(color), '<a:prstDash val="%s"/>' % dash if dash else "", he, te))
        self._append(xml)
        self.log.append(("arrow", x, y, w, h))

    def line(self, x1, y1, x2, y2, color="warm", lw=1.0, dash=None):
        self.arrow(x1, y1, x2, y2, color, lw, head=None, dash=dash)

    def image(self, path, x, y, w=None, h=None):
        pic = self.s.shapes.add_picture(path, Emu(int(x * EMU)), Emu(int(y * EMU)),
                                        Emu(int(w * EMU)) if w else None, Emu(int(h * EMU)) if h else None)
        self.log.append(("pic", x, y, pic.width / EMU, pic.height / EMU))
        return pic

    # ---- バンディットの部品（本人の資料の寸法） ----
    def player(self, x, y, size="s"):
        """プレイヤー（人型）。s: 0.26×0.66, m: 0.32×0.81。左上指定。戻り値は (中心x, 上端, 下端)。"""
        w, h = (0.26, 0.66) if size == "s" else (0.32, 0.81)
        self.image(ICON["human"], x, y, w, h)
        return x + w / 2, y, y + h

    def slot(self, x, y, size="s"):
        """腕（スロットマシン）。s: 0.86×0.46, m: 1.05×0.57。戻り値は (中心x, 上端, 下端)。"""
        w, h = (0.86, 0.46) if size == "s" else (1.05, 0.57)
        self.image(ICON["slot"], x, y, w, h)
        return x + w / 2, y, y + h

    def lock(self, x, y, h=0.64):
        """南京錠＝そのプレイヤーが腕に固定されている印。**固定されたプレイヤーの左隣**に置く
        （本人の修正 2026-10-15：腕の上に置くと腕が隠れ、誰が固定なのか読めない）。
        プレイヤー (px, py) なら lock(px - 0.6, py - 0.45)。"""
        self.image(ICON["lock"], x, y, h * 70 / 80, h)

    def collision(self, cx, cy, w=1.16, h=0.72):
        """衝突の爆発図形。**腕のアイコンには重ねない。** 2 本の矢印が腕に入る所
        （腕の下端のすぐ下、cy = 腕の下端 + 0.3 前後）に置く（本人の修正 2026-10-15）。"""
        self.shape("irregularSeal2", cx - w / 2, cy - h / 2, w, h, fill="salmon", line="red", lw=1.0)

    def collide(self, arm, players, size="s", gap=1.5):
        """「同じ腕を複数のプレイヤーが引いて衝突する」図（本人の修正版の型）。
        arm=(腕の左端 x, 上端 y)。players はプレイヤーの左端 x のリスト（離して置く。並べて平行な
        矢印にしない）。1 人目は腕の真下からまっすぐ、ほかは斜めに同じ腕の下端へ矢印を引き、
        矢印が集まる所に爆発図形を置く。戻り値は (腕の中心 x, プレイヤーの上端 y)。"""
        aw, ah = (0.86, 0.46) if size == "s" else (1.05, 0.57)
        cx, bot = arm[0] + aw / 2, arm[1] + ah
        py = bot + gap
        for i, px in enumerate(players):
            pc, _, _ = self.player(px, py)
            tx = cx + (0 if i == 0 else 0.12 * (1 if px > cx else -1))
            self.arrow(pc, py - 0.05, tx, bot + 0.05)
        self.collision(cx + 0.1, bot + 0.45, w=0.95, h=0.6)
        return cx, py

    def fieldset(self, x, y, w, h, label, body, color="green", size=18, label_size=20):
        """見出しを枠の上辺に載せた白地の箱（本人の修正 2026-10-15 の型）。
        塗りなし・色つき 2.25pt の枠、見出しは同じ色の太字を背景色の小箱で枠線の上に重ねる。
        「良い場合 / 悪い場合」のように 2〜3 個の短い結論を並べるときだけ使う。箇条書きは入れない。"""
        self.box(x, y, w, h, body, fill=None, line=color, lw=2.25, size=size, radius=0.10,
                 ins=(0.15, 0.06, 0.15, 0.06), spacing=1.5)
        lw = _cells(_plain(label)) * 0.15 * label_size / 20 + 0.5
        self._sp(x + 0.27, y - 0.30, lw, 0.54, fill="FFFEFF", wrap=False,
                 body=para(label, label_size, color, True, "ctr"), txbox=True, anchor="ctr")
        return y + h

    def points(self, x, y, w, items, color="green", heading=None, size=18, gap=12):
        """箱で囲まない箇条書き（本人の修正 2026-10-15 の型）。
        heading は 20pt 太字の色つき見出し。items は "**要点**：説明" の形の文字列で、要点を
        heading と同じ色の太字にする（行頭の • も同じ色になる）。行間 120%、段落の後に gap pt。"""
        if heading:
            y = self.text(x, y, w, heading, 20, color, bold=True) + 0.1
        def colorize(it):
            m = re.match(r"\*\*(.+?)\*\*(.*)$", it, flags=re.S)
            return "{%s:%s}%s" % (color, m.group(1), m.group(2)) if m else it
        paras = [para(colorize(it), size, "body", False, "l", "•", 1.2, after=gap) for it in items]
        est = sum(text_h(it, w - 0.3125, size, 1.2) for it in items) + gap / 72.0 * len(items)
        self._sp(x, y, w, est, body="".join(paras), txbox=True, autofit=True, tag="Points")
        return y + est

    def ring(self, cx, cy, d=0.31, color="red", lw=2.25):
        self.shape("ellipse", cx - d / 2, cy - d / 2, d, d, fill=None, line=color, lw=lw)

    def frame(self, x, y, w, h, color="red", lw=1.25, dash=None):
        self.shape("rect", x, y, w, h, fill=None, line=color, lw=lw, dash=dash)

    def code(self, x, y, w, lines, size=16, hl=None, h=None, title=None, lh=None):
        """擬似コードの箱（灰色の細枠の角丸）。キーワードは {red-n:if} のように赤の細字。
        hl={行番号: "pale-blue" | "yellow"} でその行に帯を敷く。先頭の空白2つで1段字下げ。
        折り返す行は、その行数ぶんの高さを取る（重ならない）。"""
        lh = lh or C.LINE20 * size / 20 * 1.08
        rows = []
        for ln in lines:
            depth = (len(ln) - len(ln.lstrip(" "))) // 2
            n = lines_of(ln.lstrip(" "), w - 0.2 - depth * 0.28, size)
            rows.append((ln.lstrip(" "), depth, n))
        need = sum(n for _, _, n in rows) * lh + (lh if title else 0) + 0.24
        hh = max(h or 0, need)
        self.box(x, y, w, hh, fill=None, line="line-gray", lw=1.0, radius=0.08)
        yy = y + 0.12
        if title:
            self.text(x + 0.12, yy - 0.03, w - 0.24, title, size, "warm", bold=True, h=lh)
            yy += lh
        for i, (ln, depth, n) in enumerate(rows):
            if hl and i in hl:
                self.box(x + 0.08 + depth * 0.28, yy, w - 0.16 - depth * 0.28, lh * n, fill=hl[i], prst="rect")
            self.text(x + 0.10 + depth * 0.28, yy - 0.02, w - 0.2 - depth * 0.28, ln, size,
                      h=lh * n + 0.02, anchor="ctr")
            yy += lh * n
        return y + hh

    def table(self, x, y, col_w, rows, size=16, row_h=0.40, header=True, align="ctr",
              hl_cells=None, hl_cols=None, first_col_bold=False, header_fill=None, col_align=None):
        """表。泉知成_修論中間発表スライド_完成.pptx のデザイン（2026-10-14 本人の指定）：
        縦罫なし。見出し行は塗りなし・紺 0D114A の太字・下に紺 2pt。本文の行は塗りなし・文字 404040・
        行の上下に BFBFBF 0.75pt。強調は hl_cells={(r,c): "yellow"}、hl_cols={c: "yellow"} で塗る。
        セルは markup（$数式$ 可）。"""
        nr, nc = len(rows), len(col_w)
        gf = self.s.shapes.add_table(nr, nc, Emu(int(x * EMU)), Emu(int(y * EMU)),
                                     Emu(int(sum(col_w) * EMU)), Emu(int(row_h * nr * EMU)))
        tbl = gf.table
        tblPr = tbl._tbl.find("{%s}tblPr" % NS["a"])
        for k in list(tblPr.attrib):
            del tblPr.attrib[k]
        for ch in list(tblPr):
            tblPr.remove(ch)
        for j, cw in enumerate(col_w):
            tbl.columns[j].width = Emu(int(cw * EMU))
        for i in range(nr):
            tbl.rows[i].height = Emu(int(row_h * EMU))
            for j in range(nc):
                tc = tbl.cell(i, j)._tc
                is_h = header and i == 0
                fill = header_fill if is_h else None
                if hl_cols and j in hl_cols and not is_h:
                    fill = hl_cols[j]
                if hl_cells and (i, j) in hl_cells:
                    fill = hl_cells[(i, j)]
                for old in tc.findall("{%s}txBody" % NS["a"]):
                    tc.remove(old)
                txt = rows[i][j] if j < len(rows[i]) else ""
                al = (col_align[j] if col_align else align)
                tc.insert(0, _frag('<a:txBody><a:bodyPr/><a:lstStyle/>%s</a:txBody>'
                                   % para(txt, size, "ink" if is_h else "404040",
                                          bold=is_h or (first_col_bold and j == 0), align=al)))
                tcPr = tc.find("{%s}tcPr" % NS["a"])
                if tcPr is None:
                    tcPr = etree.SubElement(tc, "{%s}tcPr" % NS["a"])
                for ch in list(tcPr):
                    tcPr.remove(ch)
                for k, v in (("anchor", "ctr"), ("marL", str(int(0.06 * EMU))), ("marR", str(int(0.06 * EMU))),
                             ("marT", "0"), ("marB", "0")):
                    tcPr.set(k, v)

                def ln(tag, w, c):
                    if w:
                        return ('<a:%s w="%d" cmpd="sng"><a:solidFill><a:srgbClr val="%s"/></a:solidFill></a:%s>'
                                % (tag, int(w * 12700), col(c), tag))
                    return '<a:%s w="12700"><a:noFill/></a:%s>' % (tag, tag)
                top = 0 if is_h else (2.0 if (header and i == 1) else 0.75)
                xml = (ln("lnL", 0, None) + ln("lnR", 0, None)
                       + ln("lnT", top, "ink" if top == 2.0 else "line-gray")
                       + ln("lnB", 2.0 if is_h else 0.75, "ink" if is_h else "line-gray"))
                xml += '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % col(fill) if fill else "<a:noFill/>"
                for e in list(_frag("<x>%s</x>" % xml)):
                    tcPr.append(e)
        el = gf._element
        xml = etree.tostring(el).decode()
        if has_math(xml):
            # PowerPoint と同じく表ごと AlternateContent で包む。代替側は数式を文字にした表
            # （Keynote は表の中の数式を描かないので、確認画像では代替側が見える）。
            parent = el.getparent()
            idx = parent.index(el)
            parent.remove(el)
            fb = re.sub(r"<a14:m\b.*?</a14:m>", _math_as_text, xml, flags=re.S)
            # Requires="a14" の接頭辞の宣言は Deck.save() で mc:Choice に足す（lxml はここで付けても省く）。
            parent.insert(idx, _frag('<mc:AlternateContent><mc:Choice Requires="a14">%s</mc:Choice>'
                                     '<mc:Fallback>%s</mc:Fallback></mc:AlternateContent>' % (xml, fb)))
        self.log.append(("table", x, y, sum(col_w), row_h * nr))
        return y + row_h * nr

    def callout(self, x, y, w, h, markup, tip, size=16, fill="white", line="ink", color="body",
                prst="wedgeRoundRectCallout"):
        """吹き出し。tip=(tx, ty) は尖った先の絶対座標。"""
        adj = {"adj1": int(((tip[0] - (x + w / 2)) / w) * 100000),
               "adj2": int(((tip[1] - (y + h / 2)) / h) * 100000)}
        if prst == "wedgeRoundRectCallout":
            adj["adj3"] = 16667
        body = "".join(para(p, size, color, align="ctr") for p in str(markup).split("\n"))
        self._sp(x, y, w, h, prst=prst, fill=fill, line=line, lw=1.0, body=body, anchor="ctr", adj=adj,
                 ins=(0.06, 0.03, 0.06, 0.03))

    def brace(self, x, y, w, h, side="left", color="warm", lw=1.5):
        """中かっこ。side="left"/"right" は (x,y,w,h) のまま、"top"/"bottom" は見た目の寸法。"""
        if side in ("left", "right"):
            self.shape("leftBrace" if side == "left" else "rightBrace", x, y, w, h, line=color, lw=lw)
        else:
            self.shape("rightBrace" if side == "bottom" else "leftBrace", x + w / 2 - h / 2, y + h / 2 - w / 2,
                       h, w, line=color, lw=lw, rot=90)

    def chevrons(self, x, y, w, h, items, fills, lines, colors=None, size=16):
        """段階の帯（先頭 homePlate、以降 chevron）。items は markup（\\n で2行）。"""
        n = len(items)
        d = h * 0.30
        cw = (w + (n - 1) * d) / n
        for i, it in enumerate(items):
            xx = x + i * (cw - d)
            body = "".join(para(p, size, (colors or ["body"] * n)[i], True, "ctr") for p in it.split("\n"))
            self._sp(xx, y, cw, h, prst="homePlate" if i == 0 else "chevron", fill=fills[i], line=lines[i],
                     lw=1.25, body=body, anchor="ctr", adj={"adj": int(d / h * 100000)},
                     ins=(d + 0.05 if i else 0.1, 0.02, d, 0.02))
        return y + h

    def tbox(self, x, y, markup, size=16, w=None):
        """"t = 3" のような時刻の小箱（細い枠）。戻り値は右端 x。"""
        # 数式は PowerPoint では見積もりより幅を取るので、狭いと "t" と "= 3" が 2 行に割れる
        w = w or max(1.05, _cells(_plain(markup)) * 0.12 * size / 20 + 0.6)
        self.box(x, y, w, 0.36, markup, fill="white", line="warm", lw=0.75, prst="rect", size=size,
                 ins=(0.03, 0.0, 0.03, 0.0))
        return x + w

    def mark(self, x, y, w, h, color="yellow"):
        """数式の項などの背面に敷くマーカー（最背面に入る）。"""
        self._sp(x, y, w, h, fill=color, back=True)

    def source(self, markup, y=7.0):
        """左下の出典（12pt）。"""
        self.text(C.L, y, 9.4, markup, 12, "warm", h=0.32)

    def quote(self, x, y, w, markup, source=None, size=18):
        """原文の引用。左に紺の縦線、淡い面、右下に出典。"""
        h = text_h(markup, w - 0.4, size) + (0.34 if source else 0) + 0.06
        self.box(x, y, w, h, fill="lead", prst="rect")
        self.box(x, y, 0.07, h, fill="ink", prst="rect")
        self.text(x + 0.2, y + 0.04, w - 0.35, markup, size, "body", italic=False,
                  h=h - (0.34 if source else 0) - 0.04)
        if source:
            self.text(x + 0.2, y + h - 0.38, w - 0.35, source, 14, "warm", align="r", h=0.34)
        return y + h

    # ---- グラフ（PowerPoint のネイティブのグラフ。画像にしない） ----
    def _style_chart(self, chart, size, legend, legend_pos):
        from pptx.enum.chart import XL_LEGEND_POSITION
        from pptx.util import Pt
        chart.font.size = Pt(size)
        chart.font.name = FONT_LATIN
        chart.font.color.rgb = RGBColor.from_string(PAL["body"])
        chart.has_title = False          # 系列が 1 つだと PowerPoint が系列名を自動のタイトルにする
        chart.has_legend = bool(legend)
        if legend:
            chart.legend.position = {"r": XL_LEGEND_POSITION.RIGHT, "b": XL_LEGEND_POSITION.BOTTOM,
                                     "t": XL_LEGEND_POSITION.TOP}[legend_pos]
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(size)

    @staticmethod
    def _axis(ax, title, rng, size, fmt=None, log=False, grid=True):
        from pptx.util import Pt
        if title:
            ax.has_title = True
            ax.axis_title.text_frame.text = title
            r = ax.axis_title.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.bold, r.font.name = Pt(size), False, FONT_LATIN
            r.font.color.rgb = RGBColor.from_string(PAL["body"])
        if rng:
            ax.minimum_scale, ax.maximum_scale = rng
        if fmt:
            ax.tick_labels.number_format, ax.tick_labels.number_format_is_linked = fmt, False
        ax.has_major_gridlines = grid
        if grid:
            ax.major_gridlines.format.line.color.rgb = RGBColor.from_string("D9D9D9")
            ax.major_gridlines.format.line.width = Pt(0.75)
        ax.format.line.color.rgb = RGBColor.from_string(PAL["warm"])
        if log:
            C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
            sc = ax._element.find("{%s}scaling" % C_NS)
            lb = etree.Element("{%s}logBase" % C_NS)
            lb.set("val", "10")
            sc.insert(0, lb)

    def xy_chart(self, x, y, w, h, series, x_title=None, y_title=None, x_range=None, y_range=None,
                 size=14, legend=True, legend_pos="r", x_fmt=None, y_fmt=None, x_log=False, y_log=False):
        """折れ線（XY）。series は (名前, xs, ys, 色, 太さpt, 破線か) のリスト。"""
        from pptx.chart.data import XyChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.util import Pt
        cd = XyChartData()
        for name, xs, ys, *_ in series:
            sr = cd.add_series(name)
            for a, b in zip(xs, ys):
                sr.add_data_point(a, b)
        gf = self.s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Emu(int(x * EMU)), Emu(int(y * EMU)),
                                     Emu(int(w * EMU)), Emu(int(h * EMU)), cd)
        ch = gf.chart
        self._style_chart(ch, size, legend, legend_pos)
        for sr, spec in zip(ch.plots[0].series, series):
            c, lw, dash = (list(spec[3:]) + [None, None, None])[:3]
            sr.smooth = False
            sr.format.line.color.rgb = RGBColor.from_string(col(c or "ink"))
            sr.format.line.width = Pt(lw or 2.25)
            if dash:
                sr.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            sr.marker.style = XL_MARKER_STYLE.NONE
        self._axis(ch.category_axis, x_title, x_range, size, x_fmt, x_log)
        self._axis(ch.value_axis, y_title, y_range, size, y_fmt, y_log)
        self.log.append(("chart", x, y, w, h))
        return ch

    def bar_chart(self, x, y, w, h, categories, series, horizontal=True, log=False, size=14,
                  legend=False, value_range=None, value_title=None, labels=True, fmt="#,##0", gap=60):
        """棒グラフ。series は (名前, 値のリスト, 色 or 色のリスト)。log=True で対数軸。"""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
        from pptx.util import Pt
        cd = CategoryChartData()
        cd.categories = categories
        for name, vals, _ in series:
            cd.add_series(name, vals)
        kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = self.s.shapes.add_chart(kind, Emu(int(x * EMU)), Emu(int(y * EMU)), Emu(int(w * EMU)),
                                     Emu(int(h * EMU)), cd)
        ch = gf.chart
        self._style_chart(ch, size, legend, "b")
        ch.plots[0].gap_width = gap
        for sr, (_, vals, c) in zip(ch.plots[0].series, series):
            sr.invert_if_negative = False
            if isinstance(c, (list, tuple)):
                for i, ci in enumerate(c):
                    pt = sr.points[i]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = RGBColor.from_string(col(ci))
            else:
                sr.format.fill.solid()
                sr.format.fill.fore_color.rgb = RGBColor.from_string(col(c))
            if labels:
                sr.data_labels.show_value = True
                sr.data_labels.number_format, sr.data_labels.number_format_is_linked = fmt, False
                sr.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
                sr.data_labels.font.size = Pt(size)
        if horizontal:
            ch.category_axis.reverse_order = True
        self._axis(ch.category_axis, None, None, size, grid=False)
        self._axis(ch.value_axis, value_title, value_range, size, fmt, log)
        self.log.append(("chart", x, y, w, h))
        return ch

    # ---- 本人の過去のスライドから図形をそのまま複製する ----
    def paste(self, src, slide_no, idx, dx=0.0, dy=0.0, retext=None, recolor=None):
        """src（.pptx）の slide_no 枚目の図形を複製する。idx は spTree の子の番号
        （scripts/xdump.py で一覧できる。数式を含む図形も番号が付く）。
        dx, dy はインチでの平行移動。retext={番号: (markup, size, color, bold, align, w)} で
        その図形の文字だけを差し替える（枠・塗りは元のまま）。size=None なら文字の書式も元のまま
        （markup は普通の文字として入る）。w を渡すと中心を保ったまま幅を変える。
        画像は自分のスライドに取り込み、図形の番号は振り直す。"""
        import copy
        import io
        prs = self.deck._src(src)
        sld = prs.slides[slide_no - 1]
        tree = sld.shapes._spTree
        kids = list(tree)
        R_NS = NS["r"]
        for i in idx:
            el = copy.deepcopy(kids[i])
            # 画像の取り込み
            for node in el.iter():
                for att in ("{%s}embed" % R_NS, "{%s}link" % R_NS, "{%s}id" % R_NS):
                    rid = node.get(att)
                    if not rid:
                        continue
                    rel = sld.part.rels.get(rid)
                    if rel is None or rel.is_external or "image" not in rel.reltype:
                        del node.attrib[att]
                        continue
                    _, new_rid = self.s.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                    node.set(att, new_rid)
            # 番号を振り直す
            for c in el.iter("{%s}cNvPr" % NS["p"]):
                c.set("id", str(self._id_after(el)))
            # 文字の差し替え
            if retext and i in retext:
                markup, size, color, bold, align, neww = (list(retext[i]) + [None] * 6)[:6]
                target = el
                if etree.QName(el).localname == "AlternateContent":
                    target = el.find("{%s}Choice" % MC)[0]
                    el = target
                if neww:
                    xf = target.find("{%s}spPr/{%s}xfrm" % (NS["p"], NS["a"]))
                    off, ext = xf.find("{%s}off" % NS["a"]), xf.find("{%s}ext" % NS["a"])
                    cx = int(off.get("x")) + int(ext.get("cx")) / 2
                    off.set("x", str(int(cx - neww * EMU / 2)))
                    ext.set("cx", str(int(neww * EMU)))
                tb = target.find("{%s}txBody" % NS["p"])
                if size is None:
                    # 書式はそのまま（最初の文字列の書式を使う）、文字だけ差し替える。数式の部分は消える。
                    first = next(tb.iter("{%s}r" % NS["a"]))
                    rp = copy.deepcopy(first.find("{%s}rPr" % NS["a"]))
                    p0 = copy.deepcopy(tb.find("{%s}p" % NS["a"]))
                    for pp in tb.findall("{%s}p" % NS["a"]):
                        tb.remove(pp)
                    for line in str(markup).split("\n"):
                        pp = copy.deepcopy(p0)
                        for ch in list(pp):
                            if etree.QName(ch).localname not in ("pPr",):
                                pp.remove(ch)
                        r = etree.SubElement(pp, "{%s}r" % NS["a"])
                        r.append(copy.deepcopy(rp))
                        etree.SubElement(r, "{%s}t" % NS["a"]).text = line
                        tb.append(pp)
                else:
                    for pp in tb.findall("{%s}p" % NS["a"]):
                        tb.remove(pp)
                    for k, line in enumerate(str(markup).split("\n")):
                        tb.append(_frag(para(line, size, color or "body", bool(bold), align or "ctr")))
            # 平行移動（最上位の xfrm だけ。グループは子の座標系ごと動く）
            tops = [el]
            if etree.QName(el).localname == "AlternateContent":
                tops = [ch[0] for ch in el if len(ch)]
            for t in tops:
                tag = etree.QName(t).localname
                if tag == "graphicFrame":
                    off = t.find("{%s}xfrm/{%s}off" % (NS["p"], NS["a"]))
                elif tag == "grpSp":
                    off = t.find("{%s}grpSpPr/{%s}xfrm/{%s}off" % (NS["p"], NS["a"], NS["a"]))
                else:
                    off = t.find("{%s}spPr/{%s}xfrm/{%s}off" % (NS["p"], NS["a"], NS["a"]))
                if off is not None:
                    off.set("x", str(int(int(off.get("x")) + dx * EMU)))
                    off.set("y", str(int(int(off.get("y")) + dy * EMU)))
            self.tree.append(el)
        return self

    def _id_after(self, el):
        self._next_id = max(getattr(self, "_next_id", 0), self._id()) + 1
        return self._next_id

    def notes(self, markup):
        """ノート。本人の形式：1 行目はタイトルの訳（太字）、空行、本文。
        行頭が「・」か「- 」の行は行頭記号つきの箇条書き、「  - 」は 2 段目。
        「【…】」で始まる行は太字の見出し。**太字**、$数式$（数式モード）が使える。"""
        tf = self.s.notes_slide.notes_text_frame
        tb = tf._txBody
        for pp in tb.findall("{%s}p" % NS["a"]):
            tb.remove(pp)
        for k, line in enumerate(str(markup).split("\n")):
            tb.append(_frag(note_para(line, first=(k == 0))))


# --- デッキ ----------------------------------------------------------------------
class Deck:
    def __init__(self, base=BASE):
        self.prs = Presentation(base)
        L = {l.name: l for l in self.prs.slide_layouts}
        self.L_cover, self.L_body, self.L_blank = L["タイトル スライド"], L["タイトルとコンテンツ"], L["白紙"]
        self.slides = []

    def _src(self, path):
        """paste() の複製元。REF の短い名前、assets/reference/ のファイル名、またはパス。"""
        if path in REF:
            path = REF[path]
        elif not os.path.exists(path) and os.path.exists(os.path.join(REF_DIR, os.path.basename(path))):
            path = os.path.join(REF_DIR, os.path.basename(path))
        if not hasattr(self, "_srcs"):
            self._srcs = {}
        if path not in self._srcs:
            self._srcs[path] = Presentation(path)
        return self._srcs[path]

    def _new(self, layout):
        s = self.prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        return s

    @staticmethod
    def _slide_number(s):
        """ページ番号（レイアウトの書式を継承。位置は 10.11, 6.95）。"""
        xml = ('<p:sp xmlns:p="%s" xmlns:a="%s"><p:nvSpPr><p:cNvPr id="90" name="Slide Number"/>'
               '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="sldNum" sz="quarter" idx="12"/></p:nvPr></p:nvSpPr>'
               '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm></p:spPr>'
               '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:fld id="{BF8D9D08-3422-A948-9BD5-C735FAB2ABFB}" type="slidenum">'
               '<a:rPr lang="en-US"/><a:t>‹#›</a:t></a:fld><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>'
               % (NS["p"], NS["a"], int(10.11 * EMU), int(6.95 * EMU), int(3.0 * EMU), int(0.40 * EMU)))
        s.shapes._spTree.append(etree.fromstring(xml))

    def cover(self, tag, title, authors, venue, date, grade, name, affiliations=(), title_size=28):
        """表紙。tag は「論文ゼミ 発表資料」など（不明なら本人に聞く）。"""
        s = self._new(self.L_cover)
        sl = Slide(self, s, 0, "cover")
        sl._sp(0.28, 0.31, _cells(tag) * 0.15 + 0.3, 0.44, line="red", lw=1.5, body=para(tag, 20, "red"),
               txbox=True, autofit=True, wrap=False)
        th = lines_of(title, 10.0, title_size) * 0.47 * title_size / 28 + 0.14
        ty = 1.62
        sl._sp(1.79, ty, 10.0, th, body=para(title, title_size, "ink", True), txbox=True, anchor="b")
        ya = ty + th + 0.06
        sl.text(1.79, ya, 10.0, authors, 20)
        yv = ya + 0.50
        sl.text(1.79, yv, 10.0, "{i:%s}" % venue, 20)
        sl.box(1.43, ty + 0.08, 0.11, yv + 0.44 - (ty + 0.08), fill="warm", prst="rect")
        sl.text(1.79, yv + 1.00, 10.0, "%s　%s  %s" % (date, grade, name), 20, "warm")
        if affiliations:
            sl.line(1.26, 5.45, 12.03, 5.45, "warm", 0.75)
            sl.text(1.26, 5.58, 11.2, "\n".join(affiliations), 14, "warm")
        self.slides.append(sl)
        return sl

    def toc(self, items, title="Table of Contents", current=None):
        """目次（24pt、行間 200%、❶〜）。current を指定すると他をグレーにする。"""
        s = self._new(self.L_body)
        sl = Slide(self, s, C.TOP_NOLEAD, "toc")
        self._title(sl, title)
        body = "".join(para("%s  %s" % ("❶❷❸❹❺❻❼❽❾❿"[i], it), 24,
                            "navy-bg" if current is None or current == i else "gray", spacing=2.0)
                       for i, it in enumerate(items))
        sl._sp(0.69, 1.18, 10.82, 0.3 + 0.8 * len(items), body=body, txbox=True, autofit=True)
        self._slide_number(s)
        self.slides.append(sl)
        return sl

    def divider(self, title):
        """セクション扉（紺地 16223B、白抜き 36pt）。"""
        s = self._new(self.L_blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(PAL["navy-bg"])
        sl = Slide(self, s, 0, "divider")
        sl._sp(0.84, 3.40, 11.69, 0.71, body=para(title, 36, "white", True), txbox=True, autofit=True)
        self._slide_number(s)
        self.slides.append(sl)
        return sl

    @staticmethod
    def _title(sl, title, part=None, size=24):
        """タイトル（24pt 太字 紺）。part="(1/3)" は同じ行の末尾に2つの空白を空けて付ける。"""
        t = "%s  %s" % (title, part) if part else title
        sl._sp(C.TITLE_X, C.TITLE_Y, 11.9, C.TITLE_H, body=para(t, size, "ink", True), txbox=True,
               anchor="ctr", tag="Title")

    def slide(self, title, lead=None, part=None, lead_fill="lead", title_size=24):
        """本文スライド。lead はそのページの主張を 1〜2 文で言い切る（20pt、帯の中）。
        part="(1/3)" で続きものの番号をタイトルの右に置く。戻り値の .top が本文の開始 y。"""
        s = self._new(self.L_body)
        sl = Slide(self, s, C.TOP_NOLEAD)
        self._title(sl, title, part, title_size)
        if lead:
            n = lines_of(lead, C.W - 0.44, 20)
            h = 0.78 if n == 1 else 1.10 if n == 2 else 1.40
            sl.box(C.L, C.LEAD_Y, C.W, h, lead, fill=lead_fill, align="l", anchor="ctr",
                   ins=(0.22, 0.05, 0.22, 0.05), radius=0.10, )
            sl.log[-1] = ("lead",) + sl.log[-1][1:]
            sl.top = C.LEAD_Y + h + 0.22
        self._slide_number(s)
        self.slides.append(sl)
        return sl

    def save(self, path):
        """保存し、mc:Choice の Requires="a14" に接頭辞 a14 の宣言を足す。
        宣言がないと MCE の規則違反で、PowerPoint は開くたびに「修復しますか？」と聞く。"""
        import io
        import zipfile
        buf = io.BytesIO()
        self.prs.save(buf)
        fix = ('<mc:Choice Requires="a14">', '<mc:Choice xmlns:a14="%s" Requires="a14">' % A14)
        with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin, \
                zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename.startswith("ppt/slides/slide") and info.filename.endswith(".xml"):
                    data = data.decode("utf-8").replace(*fix).encode("utf-8")
                if info.filename.endswith(".xml") and info.filename.startswith(
                        ("ppt/slides/slide", "ppt/slideLayouts/", "ppt/slideMasters/")):
                    # 切り替え効果は付けない（2026-10-14 本人の指定。土台の変形 morph も外す）
                    data = _strip_transitions(data.decode("utf-8")).encode("utf-8")
                zout.writestr(info, data)
        return path


if __name__ == "__main__":
    print(__doc__)
