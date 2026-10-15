import sys,re
from pptx import Presentation
from lxml import etree
E=914400
NS={"a":"http://schemas.openxmlformats.org/drawingml/2006/main","p":"http://schemas.openxmlformats.org/presentationml/2006/main","mc":"http://schemas.openxmlformats.org/markup-compatibility/2006","m":"http://schemas.openxmlformats.org/officeDocument/2006/math"}
def bbox(el):
    x=el.find(".//a:xfrm",NS) if not el.tag.endswith("graphicFrame") else el.find(".//p:xfrm",NS)
    if x is None:
        x=el.find(".//p:xfrm",NS)
    if x is None: return (-1,-1,0,0)
    off=x.find("a:off",NS); ext=x.find("a:ext",NS)
    return int(off.get("x"))/E,int(off.get("y"))/E,int(ext.get("cx"))/E,int(ext.get("cy"))/E
def text(el):
    ch=el
    if el.tag.endswith("AlternateContent"): ch=el.find("mc:Choice",NS)[0]
    return "".join(t.text or "" for t in ch.iter("{%s}t"%NS["a"],"{%s}t"%NS["m"]))
p=Presentation(sys.argv[1])
for no in map(int,sys.argv[2:]):
    tree=p.slides[no-1].shapes._spTree
    print("=== slide",no)
    for i,el in enumerate(tree):
        tag=etree.QName(el).localname
        if tag in ("nvGrpSpPr","grpSpPr","extLst"): continue
        inner=el.find("mc:Choice",NS)[0] if tag=="AlternateContent" else el
        nm=inner.find(".//p:cNvPr",NS)
        x,y,w,h=bbox(inner)
        print("%3d %-16s %-22s (%5.2f,%5.2f) %5.2fx%5.2f %s"%(i,tag+("/"+etree.QName(inner).localname if tag=="AlternateContent" else ""),(nm.get("name") if nm is not None else "")[:22],x,y,w,h,text(el)[:60]))
